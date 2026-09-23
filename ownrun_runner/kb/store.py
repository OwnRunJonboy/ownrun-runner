from __future__ import annotations

import logging
import math
import re
import sqlite3
import struct
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable

log = logging.getLogger(__name__)


@dataclass
class ParsedSection:
    text: str
    page: int | None = None
    chapter: str | None = None
    section_title: str | None = None


@dataclass
class Chunk:
    text: str
    ordinal: int
    page: int | None = None
    chapter: str | None = None
    section_title: str | None = None


SUPPORTED_FORMATS = frozenset({"pdf", "docx", "epub", "md", "txt"})

_ATX_HEADING = re.compile(r"^(#{1,6})[ \t]+(.+?)[ \t]*#*\s*$")
_SETEXT_UNDERLINE = re.compile(r"^(=+|-+)[ \t]*$")


def parse_document(path: Path, format: str) -> list[ParsedSection]:
    fmt = format.lower().strip()
    if fmt == "pdf":
        return _parse_pdf(path)
    if fmt == "docx":
        return _parse_docx(path)
    if fmt == "epub":
        return _parse_epub(path)
    if fmt == "md":
        return _parse_md(path)
    if fmt == "txt":
        return _parse_txt(path)
    raise ValueError(f"Unsupported format: {format}")


def _parse_pdf(path: Path) -> list[ParsedSection]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as e:
            raise ValueError(f"PDF is encrypted and could not be opened: {e}") from e

    sections: list[ParsedSection] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = _normalize_text(text)
        if text:
            sections.append(ParsedSection(text=text, page=i))
    return sections


def _parse_docx(path: Path) -> list[ParsedSection]:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    current_title: str | None = None
    sections: list[ParsedSection] = []

    def flush() -> None:
        nonlocal parts, current_title
        text = _normalize_text("\n".join(parts))
        if text:
            sections.append(ParsedSection(text=text, section_title=current_title))
        parts = []

    for para in doc.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        text = (para.text or "").strip()
        if not text:
            continue
        if style.startswith("Heading"):
            flush()
            current_title = text
            parts.append(text)
        else:
            parts.append(text)
    flush()
    return sections


def _parse_epub(path: Path) -> list[ParsedSection]:
    import ebooklib
    from ebooklib import epub
    from html.parser import HTMLParser

    class _Text(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self.parts: list[str] = []
            self._skip = False

        def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
            if tag in {"script", "style", "noscript"}:
                self._skip = True

        def handle_endtag(self, tag: str) -> None:
            if tag in {"script", "style", "noscript"}:
                self._skip = False
            if tag in {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6"}:
                self.parts.append("\n")

        def handle_data(self, data: str) -> None:
            if not self._skip:
                self.parts.append(data)

    book = epub.read_epub(str(path))
    sections: list[ParsedSection] = []
    chapter_i = 0
    for item in book.get_items():
        if item.get_type() != ebooklib.ITEM_DOCUMENT:
            continue
        name = item.get_name() or ""
        # Skip common nav/ncx fluff
        lower = name.lower()
        if any(x in lower for x in ("nav", "toc", "ncx", "cover")):
            continue
        raw = item.get_content()
        if not raw:
            continue
        try:
            html = raw.decode("utf-8", errors="ignore")
        except Exception:
            continue
        parser = _Text()
        try:
            parser.feed(html)
            parser.close()
        except Exception:
            continue
        text = _normalize_text("".join(parser.parts))
        if not text or len(text) < 40:
            continue
        chapter_i += 1
        title = Path(name).stem
        sections.append(ParsedSection(text=text, chapter=title or f"chapter-{chapter_i}"))
    return sections


def _parse_txt(path: Path) -> list[ParsedSection]:
    text = _normalize_text(_decode_text(path.read_bytes()))
    if not text:
        return []
    return [ParsedSection(text=text)]


def _parse_md(path: Path) -> list[ParsedSection]:
    text = _strip_markdown_front_matter(_decode_text(path.read_bytes()))
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    sections: list[ParsedSection] = []
    current_title: str | None = None
    parts: list[str] = []

    def flush() -> None:
        nonlocal parts
        body = _normalize_text("\n".join(parts))
        if body:
            sections.append(ParsedSection(text=body, section_title=current_title))
        parts = []

    i = 0
    while i < len(lines):
        line = lines[i]
        atx = _ATX_HEADING.match(line)
        if atx and not _in_fence(lines, i):
            flush()
            current_title = atx.group(2).strip()
            parts.append(current_title)
            i += 1
            continue
        if _is_setext_heading(lines, i):
            flush()
            current_title = line.strip()
            parts.append(current_title)
            i += 2
            continue
        parts.append(line)
        i += 1
    flush()
    return sections


def _decode_text(raw: bytes) -> str:
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        return raw.decode("utf-16")
    if b"\x00" in raw[:400]:
        for encoding in ("utf-16", "utf-16-le", "utf-16-be"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
    return raw.decode("utf-8-sig", errors="replace")


def _strip_markdown_front_matter(text: str) -> str:
    if not text.startswith("---"):
        return text
    lines = text.splitlines(keepends=True)
    if not lines or lines[0].strip() != "---":
        return text
    for i in range(1, len(lines)):
        if lines[i].strip() == "---":
            return "".join(lines[i + 1 :])
    return text


def _in_fence(lines: list[str], index: int) -> bool:
    fence: str | None = None
    for line in lines[:index]:
        stripped = line.lstrip()
        if stripped.startswith("```") or stripped.startswith("~~~"):
            marker = stripped[:3]
            fence = None if fence == marker else marker
    return fence is not None


def _is_setext_heading(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines) or not lines[index].strip():
        return False
    if _in_fence(lines, index):
        return False
    underline = lines[index + 1].strip()
    if len(underline) < 3 or _SETEXT_UNDERLINE.match(lines[index + 1]) is None:
        return False
    return set(underline) <= {"="} or set(underline) <= {"-"}


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_sections(
    sections: list[ParsedSection],
    *,
    chunk_size: int = 1200,
    overlap: int = 200,
    max_chunk_chars: int = 4096,
    max_chunks: int = 2000,
) -> list[Chunk]:
    chunks: list[Chunk] = []
    ordinal = 0
    chunk_size = max(200, chunk_size)
    overlap = max(0, min(overlap, chunk_size // 2))

    for section in sections:
        text = section.text.strip()
        if not text:
            continue
        if len(text) <= chunk_size:
            piece = text[:max_chunk_chars]
            if len(piece.strip()) >= 20:
                chunks.append(
                    Chunk(
                        text=piece,
                        ordinal=ordinal,
                        page=section.page,
                        chapter=section.chapter,
                        section_title=section.section_title,
                    )
                )
                ordinal += 1
                if ordinal > max_chunks:
                    raise ValueError(
                        f"Document exceeds max chunks ({max_chunks}); split the file and retry"
                    )
            continue

        start = 0
        while start < len(text):
            end = min(len(text), start + chunk_size)
            piece = text[start:end][:max_chunk_chars].strip()
            if len(piece) >= 20:
                chunks.append(
                    Chunk(
                        text=piece,
                        ordinal=ordinal,
                        page=section.page,
                        chapter=section.chapter,
                        section_title=section.section_title,
                    )
                )
                ordinal += 1
                if ordinal > max_chunks:
                    raise ValueError(
                        f"Document exceeds max chunks ({max_chunks}); split the file and retry"
                    )
            if end >= len(text):
                break
            start = max(0, end - overlap)

    return chunks


def embed_texts(
    ollama_base_url: str,
    model: str,
    texts: list[str],
    *,
    timeout: float = 120.0,
    on_progress: Callable[[int, int], None] | None = None,
) -> list[list[float]]:
    import requests

    vectors: list[list[float]] = []
    total = len(texts)
    for i, text in enumerate(texts):
        resp = requests.post(
            f"{ollama_base_url}/api/embeddings",
            json={"model": model, "prompt": text},
            timeout=timeout,
        )
        if resp.status_code >= 400:
            # Newer Ollama uses /api/embed
            resp = requests.post(
                f"{ollama_base_url}/api/embed",
                json={"model": model, "input": text},
                timeout=timeout,
            )
        resp.raise_for_status()
        data = resp.json()
        vec = data.get("embedding")
        if vec is None and isinstance(data.get("embeddings"), list) and data["embeddings"]:
            vec = data["embeddings"][0]
        if not isinstance(vec, list) or not vec:
            raise RuntimeError(f"Ollama returned no embedding for chunk {i}")
        vectors.append([float(x) for x in vec])
        if on_progress:
            on_progress(i + 1, total)
    return vectors


def _pack_vector(vec: list[float]) -> bytes:
    return struct.pack(f"<{len(vec)}f", *vec)


def _unpack_vector(blob: bytes) -> list[float]:
    n = len(blob) // 4
    return list(struct.unpack(f"<{n}f", blob))


def _cosine(a: list[float], b: list[float]) -> float:
    if not a or not b or len(a) != len(b):
        return 0.0
    dot = 0.0
    na = 0.0
    nb = 0.0
    for x, y in zip(a, b):
        dot += x * y
        na += x * x
        nb += y * y
    if na <= 0 or nb <= 0:
        return 0.0
    return dot / (math.sqrt(na) * math.sqrt(nb))


class KbVectorStore:
    """Lightweight sqlite-backed vector store (no separate daemon)."""

    def __init__(self, root: Path, embedding_model: str) -> None:
        self.root = root
        self.root.mkdir(parents=True, exist_ok=True)
        self.embedding_model = embedding_model
        self.db_path = self.root / "chunks.sqlite3"
        self._meta_path = self.root / ".embedding_model"
        self._conn = sqlite3.connect(str(self.db_path))
        self._conn.execute(
            """
            CREATE TABLE IF NOT EXISTS chunks (
                chunk_id TEXT PRIMARY KEY,
                document_id INTEGER NOT NULL,
                kb_id INTEGER NOT NULL,
                kb_slug TEXT NOT NULL,
                project TEXT,
                filename TEXT NOT NULL,
                ordinal INTEGER NOT NULL,
                page INTEGER,
                chapter TEXT,
                section_title TEXT,
                text TEXT NOT NULL,
                embedding BLOB NOT NULL,
                embedding_model TEXT NOT NULL
            )
            """
        )
        self._conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_chunks_doc ON chunks(document_id)"
        )
        self._conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_chunks_kb ON chunks(kb_id)"
        )
        self._conn.commit()
        self._check_model()

    def _check_model(self) -> None:
        if self._meta_path.is_file():
            recorded = self._meta_path.read_text(encoding="utf-8").strip()
            if recorded and recorded != self.embedding_model:
                log.warning(
                    "KB embedding model changed (%s → %s); search may require rebuild",
                    recorded,
                    self.embedding_model,
                )
        self._meta_path.write_text(self.embedding_model + "\n", encoding="utf-8")

    def recorded_model(self) -> str | None:
        if self._meta_path.is_file():
            return self._meta_path.read_text(encoding="utf-8").strip() or None
        return None

    def delete_document(self, document_id: int) -> int:
        cur = self._conn.execute(
            "DELETE FROM chunks WHERE document_id = ?", (document_id,)
        )
        self._conn.commit()
        return cur.rowcount

    def delete_kb(self, kb_id: int) -> int:
        cur = self._conn.execute("DELETE FROM chunks WHERE kb_id = ?", (kb_id,))
        self._conn.commit()
        return cur.rowcount

    def upsert_document(
        self,
        *,
        document_id: int,
        kb_id: int,
        kb_slug: str,
        project: str | None,
        filename: str,
        chunks: list[Chunk],
        embeddings: list[list[float]],
        embedding_model: str,
    ) -> int:
        if len(chunks) != len(embeddings):
            raise ValueError("chunks/embeddings length mismatch")
        self.delete_document(document_id)
        rows = []
        for ch, vec in zip(chunks, embeddings):
            chunk_id = f"{document_id}:{ch.ordinal}"
            rows.append(
                (
                    chunk_id,
                    document_id,
                    kb_id,
                    kb_slug,
                    project,
                    filename,
                    ch.ordinal,
                    ch.page,
                    ch.chapter,
                    ch.section_title,
                    ch.text,
                    _pack_vector(vec),
                    embedding_model,
                )
            )
        self._conn.executemany(
            """
            INSERT INTO chunks (
                chunk_id, document_id, kb_id, kb_slug, project, filename,
                ordinal, page, chapter, section_title, text, embedding, embedding_model
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            rows,
        )
        self._conn.commit()
        return len(rows)

    def first_chunk(self, document_id: int) -> dict[str, Any] | None:
        cur = self._conn.execute(
            """
            SELECT text, ordinal, page, chapter, section_title, kb_slug, filename
            FROM chunks WHERE document_id = ? ORDER BY ordinal ASC LIMIT 1
            """,
            (document_id,),
        )
        row = cur.fetchone()
        if not row:
            return None
        return {
            "text": row[0],
            "ordinal": row[1],
            "page": row[2],
            "chapter": row[3],
            "section_title": row[4],
            "kb": row[5],
            "filename": row[6],
        }

    def search(
        self,
        query_vec: list[float],
        *,
        allowed_kb_ids: Iterable[int],
        top_k: int = 5,
        kb_id: int | None = None,
        embedding_model: str | None = None,
    ) -> list[dict[str, Any]]:
        ids = list(allowed_kb_ids)
        if not ids:
            return []
        placeholders = ",".join("?" for _ in ids)
        params: list[Any] = list(ids)
        sql = f"SELECT * FROM chunks WHERE kb_id IN ({placeholders})"
        if kb_id is not None:
            sql += " AND kb_id = ?"
            params.append(kb_id)
        if embedding_model:
            sql += " AND embedding_model = ?"
            params.append(embedding_model)
        cur = self._conn.execute(sql, params)
        cols = [d[0] for d in cur.description]
        scored: list[tuple[float, dict[str, Any]]] = []
        for row in cur.fetchall():
            data = dict(zip(cols, row))
            vec = _unpack_vector(data["embedding"])
            score = _cosine(query_vec, vec)
            scored.append(
                (
                    score,
                    {
                        "score": round(score, 4),
                        "text": data["text"],
                        "kb": data["kb_slug"],
                        "filename": data["filename"],
                        "document_id": data["document_id"],
                        "page": data["page"],
                        "chapter": data["chapter"],
                        "ordinal": data["ordinal"],
                    },
                )
            )
        scored.sort(key=lambda x: x[0], reverse=True)
        return [item for _, item in scored[: max(1, top_k)]]

    def close(self) -> None:
        self._conn.close()
